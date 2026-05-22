"""Load 2019 presidential results from the TSE "Datos abiertos: Excel" export.

Per ADR-012 + Operational Commitments (Phase 1 backfill ordering), this is
the *first* end-to-end ETL pass. The 2019 export at
`resultados2019.tse.org.gt` is the cleanest source we have — clean enough
that the schema is iterated against it before tackling Memoria PDFs.

The loader reads a long-format XLSX with one row per
(round, municipality, candidate) tuple. Columns:

    | round | dept_code | muni_code | candidate_full_name | votes |

This is the format we *normalise to* from the published TSE export; a
maintainer who re-downloads the TSE files normalises them to this shape
via `pipeline/scrapers/tse/fixtures/build_fixture.py`. Tests run against
the committed fixture at `fixtures/tse_2019_presidential.xlsx`; no
live HTTP in tests.

Idempotent: re-running yields zero net presidential_results changes.
"""

from __future__ import annotations

import logging
from collections import defaultdict
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any, cast

from openpyxl import load_workbook

logger = logging.getLogger(__name__)

SCRAPE_RUN_SOURCE = "tse_excel_2019_presidential"
CYCLE = 2019
ROUNDS = (1, 2)

FIXTURE_PATH = Path(__file__).parent / "fixtures" / "tse_2019_presidential.xlsx"


@dataclass(frozen=True)
class CandidateResultRow:
    round: int
    dept_code: str
    muni_code: str
    candidate_full_name: str
    votes: int


@dataclass(frozen=True)
class LoadResult:
    round1_rows: int
    round2_rows: int
    round1_total: int
    round2_total: int
    candidates_created: int


class TotalMismatchError(RuntimeError):
    """Raised when the loaded round's national total drifts from the
    asserted TSE published total by more than the ±1-vote tolerance."""


def parse_xlsx(path: Path = FIXTURE_PATH) -> list[CandidateResultRow]:
    """Parse the long-format XLSX into typed rows. Header row required."""
    if not path.exists():
        raise FileNotFoundError(f"TSE 2019 fixture not found at {path}")

    wb = load_workbook(filename=path, read_only=True, data_only=True)
    if "results" not in wb.sheetnames:
        raise ValueError(
            f"workbook at {path} missing required 'results' sheet "
            f"(found: {wb.sheetnames})"
        )
    ws = wb["results"]

    rows: list[CandidateResultRow] = []
    expected_header = ("round", "dept_code", "muni_code", "candidate_full_name", "votes")
    header_seen = False
    for raw in ws.iter_rows(values_only=True):
        if not header_seen:
            header = tuple(str(c).strip() if c is not None else "" for c in raw[:5])
            if header != expected_header:
                raise ValueError(
                    f"workbook at {path} has unexpected header {header}; "
                    f"expected {expected_header}"
                )
            header_seen = True
            continue
        if raw is None or all(c is None for c in raw):
            continue
        rnd, dept, muni, cand, votes = raw[:5]
        if rnd is None or dept is None or muni is None or cand is None or votes is None:
            continue
        rows.append(
            CandidateResultRow(
                round=int(str(rnd)),
                dept_code=str(dept).strip(),
                muni_code=str(muni).strip(),
                candidate_full_name=str(cand).strip(),
                votes=int(str(votes)),
            )
        )
    wb.close()
    return rows


def ensure_election(conn: Any, cycle: int, round_: int) -> int:
    """Return `election_id` for (cycle, round); insert it if absent.

    Uses the UNIQUE(cycle, round) index to make the insert idempotent.
    """
    with conn.cursor() as cur:
        cur.execute(
            """
            INSERT INTO elections (cycle, round)
            VALUES (%s, %s)
            ON CONFLICT (cycle, round) DO NOTHING
            """,
            (cycle, round_),
        )
        cur.execute(
            "SELECT election_id FROM elections WHERE cycle = %s AND round = %s",
            (cycle, round_),
        )
        row = cur.fetchone()
    if row is None:
        raise RuntimeError(
            f"ensure_election: failed to materialise (cycle={cycle}, round={round_})"
        )
    return cast(int, row[0])


def ensure_candidate(conn: Any, full_name: str) -> tuple[int, bool]:
    """Return (candidate_id, created). Soft-key on `full_name`.

    There is no UNIQUE constraint on `candidates.full_name` (see
    migration 0002), so we lookup-then-insert inside the caller's
    transaction. Two concurrent loaders against the same DB could race
    here; we don't run two concurrent loaders for this source.
    """
    with conn.cursor() as cur:
        cur.execute(
            "SELECT candidate_id FROM candidates WHERE full_name = %s",
            (full_name,),
        )
        row = cur.fetchone()
        if row is not None:
            return cast(int, row[0]), False
        cur.execute(
            "INSERT INTO candidates (full_name) VALUES (%s) RETURNING candidate_id",
            (full_name,),
        )
        new = cur.fetchone()
    if new is None:
        raise RuntimeError(f"ensure_candidate: insert returned no row for {full_name!r}")
    return cast(int, new[0]), True


def resolve_geography(conn: Any, level: str, code: str) -> int:
    """Lookup `geography_id` by (level, code). Raises if absent.

    Per ADR-012 the loader writes at municipality grain, so the typical
    call is `resolve_geography(conn, "municipality", "0101")`. The
    geographies seed (issue #2) must have been applied first.
    """
    with conn.cursor() as cur:
        cur.execute(
            "SELECT geography_id FROM geographies WHERE level = %s AND code = %s",
            (level, code),
        )
        row = cur.fetchone()
    if row is None:
        raise KeyError(f"unknown geography ({level!r}, {code!r}) — seed geographies first")
    return cast(int, row[0])


def stamp_scrape_run(
    conn: Any,
    source: str = SCRAPE_RUN_SOURCE,
    *,
    success: bool,
    error_message: str | None = None,
) -> None:
    """UPSERT a row into `scrape_runs` for this loader's `source` key."""
    with conn.cursor() as cur:
        cur.execute(
            """
            INSERT INTO scrape_runs (source, last_run_at, success, error_message)
            VALUES (%s, now(), %s, %s)
            ON CONFLICT (source) DO UPDATE
              SET last_run_at = EXCLUDED.last_run_at,
                  success = EXCLUDED.success,
                  error_message = EXCLUDED.error_message
            """,
            (source, success, error_message),
        )


def load(
    conn: Any,
    xlsx_path: Path = FIXTURE_PATH,
    *,
    expected_round1_total: int | None = None,
    expected_round2_total: int | None = None,
    tolerance_votes: int = 1,
) -> LoadResult:
    """End-to-end: parse XLSX, upsert dims, bulk-upsert `presidential_results`,
    assert totals, stamp `scrape_runs`.

    `expected_round{1,2}_total` are the TSE published national totals.
    If supplied, each round's loaded sum must match to within
    `tolerance_votes` (default 1 vote, per the issue body).

    Idempotent: a second call with the same XLSX yields zero net changes
    in `presidential_results` (the ON CONFLICT updates rewrite votes to
    the same value).
    """
    rows = parse_xlsx(xlsx_path)
    if not rows:
        _safe_stamp(conn, success=False, error_message="empty fixture")
        raise ValueError(f"no rows parsed from {xlsx_path}")

    by_round: dict[int, list[CandidateResultRow]] = defaultdict(list)
    for r in rows:
        by_round[r.round].append(r)

    for rnd in by_round:
        if rnd not in ROUNDS:
            _safe_stamp(conn, success=False, error_message=f"unexpected round {rnd}")
            raise ValueError(f"row references round={rnd}; expected one of {ROUNDS}")

    try:
        election_ids = {rnd: ensure_election(conn, CYCLE, rnd) for rnd in ROUNDS}

        candidates_created = 0
        candidate_ids: dict[str, int] = {}
        all_names = {r.candidate_full_name for r in rows}
        for name in sorted(all_names):
            cid, created = ensure_candidate(conn, name)
            candidate_ids[name] = cid
            if created:
                candidates_created += 1

        geography_ids: dict[str, int] = {}
        all_munis = {r.muni_code for r in rows}
        for code in sorted(all_munis):
            geography_ids[code] = resolve_geography(conn, "municipality", code)

        round_rowcounts: dict[int, int] = {1: 0, 2: 0}
        round_totals: dict[int, int] = {1: 0, 2: 0}
        with conn.cursor() as cur:
            for r in rows:
                cur.execute(
                    """
                    INSERT INTO presidential_results
                        (election_id, candidate_id, geography_id, votes)
                    VALUES (%s, %s, %s, %s)
                    ON CONFLICT (election_id, candidate_id, geography_id) DO UPDATE
                      SET votes = EXCLUDED.votes
                    """,
                    (
                        election_ids[r.round],
                        candidate_ids[r.candidate_full_name],
                        geography_ids[r.muni_code],
                        r.votes,
                    ),
                )
                round_rowcounts[r.round] += 1
                round_totals[r.round] += r.votes

        _assert_total(1, round_totals[1], expected_round1_total, tolerance_votes)
        _assert_total(2, round_totals[2], expected_round2_total, tolerance_votes)
    except Exception as exc:
        conn.rollback()
        _safe_stamp(conn, success=False, error_message=str(exc)[:500])
        raise
    else:
        conn.commit()
        stamp_scrape_run(conn, success=True, error_message=None)
        conn.commit()

    return LoadResult(
        round1_rows=round_rowcounts[1],
        round2_rows=round_rowcounts[2],
        round1_total=round_totals[1],
        round2_total=round_totals[2],
        candidates_created=candidates_created,
    )


def _assert_total(
    round_: int,
    observed: int,
    expected: int | None,
    tolerance: int,
) -> None:
    if expected is None:
        return
    if abs(observed - expected) > tolerance:
        raise TotalMismatchError(
            f"round {round_} national total mismatch: observed={observed} "
            f"expected={expected} tolerance=±{tolerance}"
        )


def _safe_stamp(conn: Any, *, success: bool, error_message: str | None) -> None:
    """Stamp `scrape_runs` in a *separate* transaction so a load error
    on the dim/fact side does not roll back the audit row.
    """
    try:
        stamp_scrape_run(conn, success=success, error_message=error_message)
        conn.commit()
    except Exception:
        logger.exception("stamp_scrape_run failed")
        conn.rollback()


__all__ = [
    "CYCLE",
    "FIXTURE_PATH",
    "SCRAPE_RUN_SOURCE",
    "CandidateResultRow",
    "LoadResult",
    "TotalMismatchError",
    "ensure_candidate",
    "ensure_election",
    "load",
    "parse_xlsx",
    "resolve_geography",
    "stamp_scrape_run",
]


def _iter_test_payload() -> Iterable[CandidateResultRow]:
    """Diagnostic helper — yields the parsed default fixture rows."""
    yield from parse_xlsx(FIXTURE_PATH)
